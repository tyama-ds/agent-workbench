"""Real wire adapters and files; failed batches must remain safe to continue."""
import hashlib
import json

import openpyxl
import pytest
from aiohttp import web

from workbench.engine import Engine
from workbench.harness import ToolExecutor
from workbench.redaction import RedactionCapacityError, REDACTION_CAPACITY_MESSAGE
from test_engine import ScriptClient, configured, reply, settled, start
from test_engine_integration import (
    api_server, configured_settings, local_reply, openai_reply, anthropic_reply,
)


WIRE_REPLY = {'local': local_reply, 'openai': openai_reply, 'anthropic': anthropic_reply}
FOLLOW_UP = 'Do not retry any tool. Report the partial outcome from the existing results.'


def tool_results(kind, payload, ids, *, follow_up=None):
    """Assert exact pairing, order and the position of any explicit continuation."""
    if kind == 'openai':
        items = payload['input']
        calls = [item['call_id'] for item in items if item.get('type') == 'function_call']
        results = [item for item in items if item.get('type') == 'function_call_output']
        assert calls == ids and [item['call_id'] for item in results] == ids
        first = next(i for i, item in enumerate(items) if item.get('type') == 'function_call_output')
        assert items[first:first + len(ids)] == results
        if follow_up is not None:
            assert items[first + len(ids):] == [{'role': 'user', 'content': follow_up}]
        else:
            assert first + len(ids) == len(items)
        return [json.loads(item['output']) for item in results]
    if kind == 'anthropic':
        messages = payload['messages']
        calls = [block['id'] for message in messages if message['role'] == 'assistant'
                 for block in message['content'] if block['type'] == 'tool_use']
        assert calls == ids and messages[-2]['role'] == 'assistant' and messages[-1]['role'] == 'user'
        blocks = messages[-1]['content']
        results = [block for block in blocks if block['type'] == 'tool_result']
        assert [item['tool_use_id'] for item in results] == ids
        assert blocks[:len(ids)] == results
        assert blocks[len(ids):] == ([] if follow_up is None else [{'type': 'text', 'text': follow_up}])
        return [json.loads(item['content']) for item in results]
    messages = payload['messages']
    calls = [call['id'] for message in messages if message['role'] == 'assistant'
             for call in message.get('tool_calls', [])]
    results = [message for message in messages if message['role'] == 'tool']
    assert calls == ids and [item['tool_call_id'] for item in results] == ids
    first = next(i for i, item in enumerate(messages) if item['role'] == 'tool')
    assert messages[first:first + len(ids)] == results
    if follow_up is not None:
        assert messages[first + len(ids):] == [{'role': 'user', 'content': follow_up}]
    else:
        assert first + len(ids) == len(messages)
    return [json.loads(item['content']) for item in results]


@pytest.mark.parametrize('kind', WIRE_REPLY)
async def test_merged_cell_rejection_preserves_workbook_and_returns_a_tool_result(tmp_path, kind):
    requests, errors = [], []
    before = None

    async def handler(request):
        try:
            body = await request.json()
            requests.append(body)
            if len(requests) == 1:
                return WIRE_REPLY[kind]([
                    ('saved', 'write_text', {'path': 'before.txt', 'text': 'Saved first', 'expected_sha256': 'missing'}),
                    ('merged', 'xlsx_write', {'path': 'merged.xlsx', 'expected_sha256': hashlib.sha256(before).hexdigest(),
                        'cells': [{'sheet': 'Sheet', 'cell': 'A2', 'value': 'Must not be saved'},
                                  {'sheet': 'Sheet', 'cell': 'B1', 'value': 'Rejected'}]}),
                    ('read', 'xlsx_read', {'path': 'merged.xlsx', 'range': 'A1:B2'}),
                ])
            assert len(requests) == 2
            results = tool_results(kind, body, ['saved', 'merged', 'read'])
            assert results[0]['operation'] == 'created'
            assert results[1]['ok'] is False and 'non-anchor merged cell' in results[1]['error']
            assert results[2]['cells'] == [{'cell': 'A1', 'value': 'Preserve', 'formula': False}]
            return WIRE_REPLY[kind](text='The merged-cell edit was rejected; the workbook is unchanged.')
        except Exception as exc:
            errors.append(repr(exc))
            return web.json_response({'error': 'Synthetic fixture assertion failed'}, status=500)

    async with api_server(handler) as endpoint:
        settings, work = configured_settings(tmp_path, endpoint)
        workbook = openpyxl.Workbook()
        workbook.active['A1'] = 'Preserve'
        workbook.active.merge_cells('A1:B1')
        workbook.save(work / 'merged.xlsx')
        workbook.close()
        before = (work / 'merged.xlsx').read_bytes()
        engine = Engine(settings)
        try:
            _, agent = await start(engine, pm_profile=kind, max_workers=0)
            await settled(engine)
            assert not errors and len(requests) == 2
            assert agent.status == 'done' and not agent.last_error
            assert (work / 'merged.xlsx').read_bytes() == before
            assert len(agent.output_receipts) == 1 and agent.output_receipts[0]['path'] == str(work / 'before.txt')
            assert len(agent.results) == 1
            # An explicitly requested anchor-cell write remains supported. No
            # unmerge, redirected write, or automatic anchor substitution occurs.
            run = engine.runs[agent.run_id]
            await engine.execute_tool(run, agent, 'xlsx_write', {'path': 'merged.xlsx',
                'expected_sha256': hashlib.sha256(before).hexdigest(),
                'cells': [{'sheet': 'Sheet', 'cell': 'A1', 'value': 'Explicit anchor'}]})
            check = openpyxl.load_workbook(work / 'merged.xlsx')
            try:
                assert check.active['A1'].value == 'Explicit anchor'
                assert str(check.active.merged_cells) == 'A1:B1'
                assert check.active['A2'].value is None
            finally:
                check.close()
        finally:
            await engine.close()


@pytest.mark.parametrize('kind', WIRE_REPLY)
@pytest.mark.parametrize('failure_point', ['before_execution', 'after_commit', 'after_receipt',
    'set_result', 'cyclic_result', 'list_result', 'get_failure', 'nonfinite_result', 'tool_log'])
async def test_unexpected_failure_closes_batch_without_replay_or_false_rollback(tmp_path, monkeypatch, kind, failure_point):
    requests, errors, dispatches = [], [], []
    original = ToolExecutor._execute

    def fault(self, name, args):
        dispatches.append(args.get('path'))
        if args.get('path') == 'uncertain.txt' and failure_point == 'before_execution':
            raise AttributeError('PRIVATE-FAULT-DETAIL before mutation')
        result = original(self, name, args)
        if args.get('path') == 'uncertain.txt' and failure_point == 'after_commit':
            raise RuntimeError('PRIVATE-FAULT-DETAIL after actual mutation')
        return result

    monkeypatch.setattr(ToolExecutor, '_execute', fault)

    async def handler(request):
        try:
            body = await request.json()
            requests.append(body)
            if len(requests) == 1:
                return WIRE_REPLY[kind]([(ident, 'write_text', {'path': path, 'text': ident, 'expected_sha256': 'missing'})
                                        for ident, path in [('saved', 'before.txt'), ('uncertain', 'uncertain.txt'), ('skipped', 'after.txt')]])
            assert len(requests) == 2
            results = tool_results(kind, body, ['saved', 'uncertain', 'skipped'], follow_up=FOLLOW_UP)
            assert results[0]['operation'] == 'created'
            assert results[1]['ok'] is False and 'completion was not confirmed' in results[1]['error']
            assert results[2] == {'ok': False, 'error': 'Not executed: an earlier tool raised an unexpected error.'}
            assert 'PRIVATE-FAULT-DETAIL' not in json.dumps(body)
            return WIRE_REPLY[kind](text='The earlier write succeeded. The interrupted operation needs inspection.')
        except Exception as exc:
            errors.append(repr(exc))
            return web.json_response({'error': 'Synthetic fixture assertion failed'}, status=500)

    async with api_server(handler) as endpoint:
        settings, work = configured_settings(tmp_path, endpoint)
        engine = Engine(settings)
        execute = engine.execute_tool
        log = engine.log

        class BrokenResult(dict):
            def get(self, *args):
                raise RuntimeError('PRIVATE-FAULT-DETAIL while inspecting result')

        async def after_receipt(run, agent, name, args):
            result = await execute(run, agent, name, args)
            if args.get('path') == 'uncertain.txt':
                if failure_point == 'after_receipt':
                    raise IndexError('PRIVATE-FAULT-DETAIL after receipt')
                if failure_point == 'set_result':
                    return {'invalid': {'PRIVATE-FAULT-DETAIL'}}
                if failure_point == 'cyclic_result':
                    result['cycle'] = result
                if failure_point == 'list_result':
                    return [result]
                if failure_point == 'get_failure':
                    return BrokenResult(result)
                if failure_point == 'nonfinite_result':
                    result['invalid'] = float('nan')
            return result

        def fail_tool_log(agent, category, value, thinking=''):
            if category == 'tool' and 'uncertain.txt' in value and failure_point == 'tool_log':
                raise RuntimeError('PRIVATE-FAULT-DETAIL while logging')
            return log(agent, category, value, thinking)

        monkeypatch.setattr(engine, 'execute_tool', after_receipt)
        monkeypatch.setattr(engine, 'log', fail_tool_log)
        try:
            run, agent = await start(engine, pm_profile=kind, max_workers=0)
            await settled(engine)
            assert len(requests) == 1 and not errors
            assert agent.status == 'error' and run['status'] == 'waiting'
            assert '処理の完了は確認できません' in agent.last_error
            assert run['tool_calls'] == 2 and run['model_calls'] == 1
            results = [item for item in agent.conversation if item['role'] == 'tool']
            assert [item['tool_call_id'] for item in results] == ['saved', 'uncertain', 'skipped']
            assert (work / 'before.txt').read_text() == 'saved'
            assert (work / 'uncertain.txt').exists() == (failure_point != 'before_execution')
            assert not (work / 'after.txt').exists()
            receipts = list(agent.output_receipts)
            assert len(receipts) == (1 if failure_point in {'before_execution', 'after_commit'} else 2)
            assert not agent.results and engine.reservations == {}
            assert 'PRIVATE-FAULT-DETAIL' not in json.dumps(engine.snapshot())
            assert engine.message_eligibility(agent)['allowed']

            await engine.human_message(agent.id, FOLLOW_UP)
            await settled(engine)
            assert not errors and len(requests) == 2
            assert agent.status == 'done' and run['status'] == 'done'
            assert dispatches == ['before.txt', 'uncertain.txt']
            assert list(agent.output_receipts) == receipts and not (work / 'after.txt').exists()
            assert run['tool_calls'] == 2 and run['model_calls'] == 2
            assert len(agent.results) == 1
        finally:
            await engine.close()


@pytest.mark.parametrize('kind', WIRE_REPLY)
@pytest.mark.parametrize('invalid', ['malformed', 'duplicate', 'unknown'])
async def test_invalid_provider_batch_cannot_execute_an_earlier_valid_call(tmp_path, kind, invalid):
    requests = []

    async def handler(request):
        requests.append(await request.json())
        response = WIRE_REPLY[kind]([
            ('first', 'write_text', {'path': 'must-not-exist.txt', 'text': 'No', 'expected_sha256': 'missing'}),
            ('second', 'list_team', {}),
        ])
        body = json.loads(response.body)
        if kind == 'local':
            second = body['choices'][0]['message']['tool_calls'][1]
            if invalid == 'malformed': second['function']['arguments'] = 'not JSON'
            elif invalid == 'duplicate': second['id'] = 'first'
            else: second['function']['name'] = 'unlisted_tool'
        elif kind == 'openai':
            second = body['output'][1]
            if invalid == 'malformed': second['arguments'] = 'not JSON'
            elif invalid == 'duplicate': second['call_id'] = 'first'
            else: second['name'] = 'unlisted_tool'
        else:
            second = body['content'][1]
            if invalid == 'malformed': second['input'] = ['not an object']
            elif invalid == 'duplicate': second['id'] = 'first'
            else: second['name'] = 'unlisted_tool'
        return web.json_response(body)

    async with api_server(handler) as endpoint:
        settings, work = configured_settings(tmp_path, endpoint)
        engine = Engine(settings)
        try:
            run, agent = await start(engine, pm_profile=kind, max_workers=0)
            await settled(engine)
            assert len(requests) == 1 and agent.status == 'error'
            assert run['tool_calls'] == 0 and not list(work.iterdir())
            assert not agent.output_receipts and not agent.results
            assert all(item['role'] == 'user' for item in agent.conversation)
        finally:
            await engine.close()


@pytest.mark.parametrize('stage', ['execution', 'tool_log'])
async def test_redaction_capacity_still_uses_fixed_fail_closed_path(tmp_path, monkeypatch, stage):
    engine, work = configured(tmp_path, ScriptClient([reply(calls=[('list_team', {}),
        ('write_text', {'path': 'later.txt', 'text': 'No', 'expected_sha256': 'missing'})])]))

    async def saturated(*args):
        raise RedactionCapacityError()

    if stage == 'execution':
        monkeypatch.setattr(engine, 'execute_tool', saturated)
    else:
        log = engine.log
        def saturated_log(agent, category, value, thinking=''):
            if category == 'tool':
                raise RedactionCapacityError()
            return log(agent, category, value, thinking)
        monkeypatch.setattr(engine, 'log', saturated_log)
    try:
        _, agent = await start(engine)
        await settled(engine)
        assert agent.status == 'error' and agent.last_error == REDACTION_CAPACITY_MESSAGE
        assert not (work / 'later.txt').exists()
        assert not agent.output_receipts
        assert all(item['role'] != 'tool' for item in agent.conversation)
        assert all('Unexpected tool error' not in item['text'] for item in agent.logs)
        assert len(engine.client.calls) == 1
    finally:
        await engine.close()


async def test_unexpected_failure_at_tool_limit_keeps_ambiguous_outcome_visible(tmp_path, monkeypatch):
    client = ScriptClient([reply(calls=[
        ('write_text', {'path': 'before.txt', 'text': 'Saved', 'expected_sha256': 'missing'}),
        ('write_text', {'path': 'uncertain.txt', 'text': 'Committed', 'expected_sha256': 'missing'}),
        ('write_text', {'path': 'later.txt', 'text': 'No', 'expected_sha256': 'missing'}),
    ]), reply('Text-only continuation')])
    engine, work = configured(tmp_path, client)
    engine.settings.value['limits']['max_tool_calls'] = 2
    execute = engine.execute_tool

    async def fault(run, agent, name, args):
        result = await execute(run, agent, name, args)
        if args.get('path') == 'uncertain.txt':
            raise RuntimeError('PRIVATE-FAULT-DETAIL')
        return result

    monkeypatch.setattr(engine, 'execute_tool', fault)
    try:
        run, agent = await start(engine)
        await settled(engine)
        assert '処理の完了は確認できません' in agent.last_error
        results = [json.loads(item['content']) for item in agent.conversation if item['role'] == 'tool']
        assert len(results) == 3 and 'earlier tool raised an unexpected error' in results[-1]['error']
        assert run['tool_calls'] == 2 and len(agent.output_receipts) == 2
        assert (work / 'uncertain.txt').read_text() == 'Committed' and not (work / 'later.txt').exists()
        await engine.human_message(agent.id, FOLLOW_UP)
        await settled(engine)
        assert agent.status == 'done' and run['tool_calls'] == 2
    finally:
        await engine.close()
