"""Python transfers, fail-closed behavior, and opt-in real Windows isolation."""
import asyncio
import hashlib
import json
import os
from pathlib import Path
import sys
import threading
import time
import uuid

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from workbench.harness import HarnessError, PathHarness
from workbench import python_runner as runner
from workbench.windows_sandbox import SandboxError
from workbench.windows_sandbox import _prune_private_links
from workbench.windows_sandbox import _capability_names


@pytest.fixture
def setup(tmp_path, monkeypatch):
    source, target, denied = (tmp_path / part for part in ('read', 'write', 'denied'))
    for folder in (source, target, denied):
        folder.mkdir()
    harness = PathHarness([str(source), str(denied)], [str(target)], [str(denied)])
    config = {'enabled': True, 'executable': str(Path(sys.base_prefix) / 'python.exe'),
              'library_paths': [], 'timeout_seconds': 10, 'max_memory_mb': 256, 'max_output_bytes': 4096}
    monkeypatch.setattr(runner, '_runtime', lambda *_: (Path(config['executable']), [], []))
    return harness, config, source, target, denied


def execute(setup, code='print(1)', **kw):
    harness, config, *_ = setup
    return asyncio.run(runner.execute_python(config, harness, code, **kw))


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_inputs_outputs_and_each_receipt(setup, monkeypatch):
    harness, config, source, target, denied = setup
    (source / 'data.txt').write_text('input', encoding='utf-8')
    saved = []
    def child(command, workspace, access, **limits):
        assert command[1:7] == ['-I', '-S', '-B', '-u', '-X', 'utf8']
        assert (workspace / 'in' / 'data.txt').read_text() == 'input'
        assert limits['memory_mb'] == 256
        assert str(target) not in (workspace / '.awb_bootstrap.py').read_text()
        (workspace / 'out' / 'result.txt').write_text('produced', encoding='utf-8')
        return {'exit_code': 0, 'termination_reason': None, 'stdout': 'ok', 'stderr': ''}
    monkeypatch.setattr(runner, 'run_isolated', child)
    result = execute(setup, inputs=[{'path': str(source / 'data.txt'), 'name': 'in/data.txt'}],
                     outputs=[{'path': str(target / 'result.txt'), 'name': 'out/result.txt', 'expected_sha256': 'missing'}],
                     on_complete=saved.append)
    assert (target / 'result.txt').read_text() == 'produced'
    assert saved == result['outputs'] and len(saved) == 1
    assert (source / 'data.txt').read_text() == 'input'


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_denied_input_or_output_fails_before_launch(setup, monkeypatch):
    harness, config, source, target, denied = setup
    (denied / 'secret.txt').write_text('private')
    monkeypatch.setattr(runner, 'run_isolated', lambda *a, **k: pytest.fail('Must not launch'))
    with pytest.raises(HarnessError, match='denied'):
        execute(setup, inputs=[{'path': str(denied / 'secret.txt'), 'name': 'x'}])
    with pytest.raises(HarnessError, match='denied'):
        execute(setup, outputs=[{'path': str(denied / 'out.txt'), 'name': 'x', 'expected_sha256': 'missing'}])


@pytest.mark.parametrize('name', ['../escape', '/absolute', 'C:\\absolute', 'a:stream', '.awb_script.py', 'a/../x'])
def test_mapping_names_reject_escapes(name):
    with pytest.raises(HarnessError):
        runner._mappings([{'path': 'ignored', 'name': name}])


def test_private_cleanup_keeps_files_and_removes_only_hardlink(tmp_path):
    private, outside = tmp_path / 'private', tmp_path / 'outside'
    private.mkdir()
    outside.write_text('original')
    (private / 'normal.txt').write_text('keep')
    os.link(outside, private / 'link.txt')
    _prune_private_links(private)
    assert (private / 'normal.txt').read_text() == 'keep'
    assert not (private / 'link.txt').exists()
    assert outside.read_text() == 'original' and outside.stat().st_nlink == 1


def test_runtime_cannot_include_denied_data_or_links(tmp_path):
    runtime, denied = tmp_path/'runtime', tmp_path/'runtime'/'private'
    denied.mkdir(parents=True)
    harness = PathHarness([], [], [str(denied)])
    with pytest.raises(HarnessError, match='denied'):
        runner._runtime_path(str(runtime), harness, directory=True)
    original = tmp_path/'source.txt'
    original.write_text('data')
    os.link(original, runtime/'alias.txt')
    with pytest.raises(HarnessError, match='hard link'):
        runner._safe_tree(runtime)


def test_disabled_execution_does_not_probe_or_launch(setup, monkeypatch):
    setup[1]['enabled'] = False
    monkeypatch.setattr(runner, 'run_isolated', lambda *a, **k: pytest.fail('Must not launch'))
    with pytest.raises(HarnessError, match='disabled'):
        execute(setup)


@pytest.mark.parametrize('enabled,expected', [
    (False, ('registryRead',)),
    (True, ('registryRead', 'internetClient', 'privateNetworkClientServer')),
])
def test_network_capabilities_are_minimal(enabled, expected):
    assert _capability_names(enabled) == expected


@pytest.mark.parametrize('value', [None, 0, 1, 'false', 'true'])
def test_network_requires_boolean_before_native_launch(value, monkeypatch):
    from workbench import windows_sandbox as sandbox
    monkeypatch.setattr(sandbox, '_Native', lambda: pytest.fail('Must not touch Windows'))
    with pytest.raises(SandboxError, match='boolean'):
        sandbox.run_isolated([], Path('.'), [], timeout=1, memory_mb=128,
                             output_limit=1024, workspace_limit=1024, network_enabled=value)


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
@pytest.mark.parametrize('enabled', [None, False, True])
def test_network_setting_reaches_native_runner(setup, monkeypatch, enabled):
    if enabled is not None:
        setup[1]['network_enabled'] = enabled
    calls = []
    def child(*args, **limits):
        calls.append(limits['network_enabled'])
        return {'exit_code': 0, 'termination_reason': None, 'stdout': '', 'stderr': ''}
    monkeypatch.setattr(runner, 'run_isolated', child)
    execute(setup)
    assert calls == [enabled is True]


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_invalid_network_setting_fails_before_runtime_probe(setup, monkeypatch):
    setup[1]['network_enabled'] = 'true'
    monkeypatch.setattr(runner, '_runtime', lambda *a: pytest.fail('Must not probe runtime'))
    monkeypatch.setattr(runner, 'run_isolated', lambda *a, **k: pytest.fail('Must not launch'))
    with pytest.raises(HarnessError, match='boolean'):
        execute(setup)


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
@pytest.mark.parametrize('exit_code,reason', [(1, None), (0, 'timeout'), (0, 'output_limit')])
def test_failed_execution_never_publishes(setup, monkeypatch, exit_code, reason):
    target = setup[3] / 'result.txt'
    def child(command, workspace, *a, **k):
        (workspace / 'out.txt').write_text('must not publish')
        return {'exit_code': exit_code, 'termination_reason': reason, 'stdout': '', 'stderr': ''}
    monkeypatch.setattr(runner, 'run_isolated', child)
    result = execute(setup, outputs=[{'path': str(target), 'name': 'out.txt', 'expected_sha256': 'missing'}])
    assert not target.exists() and result['outputs'] == []


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_fail_closed_when_lpac_unavailable(setup, monkeypatch):
    def failed(*args, **kw):
        raise SandboxError('CreateProcessW(LPAC) failed')
    monkeypatch.setattr(runner, 'run_isolated', failed)
    with pytest.raises(HarnessError, match='LPAC'):
        execute(setup)


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_unconfirmed_child_exit_preserves_private_storage(setup, monkeypatch):
    seen = []
    def failed(command, workspace, *args, **kw):
        seen.append(workspace)
        raise SandboxError('synthetic child exit not confirmed', preserve_private=True)
    monkeypatch.setattr(runner, 'run_isolated', failed)
    try:
        with pytest.raises(HarnessError, match='private storage retained'):
            execute(setup)
        assert (seen[0]/'.awb_script.py').exists()
    finally:
        # This synthetic child never starts; test-owned storage can be removed.
        import shutil
        for path in seen:
            assert path.resolve() == path and path.name.startswith('awb-python-')
            shutil.rmtree(path)


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_partial_publication_reports_only_actual_receipts(setup, monkeypatch):
    harness, config, source, target, denied = setup
    def child(command, workspace, *a, **k):
        for name in ('one', 'two'):
            (workspace / name).write_text(name)
        return {'exit_code': 0, 'termination_reason': None, 'stdout': '', 'stderr': ''}
    monkeypatch.setattr(runner, 'run_isolated', child)
    saved = []
    original_write = harness.write_bytes
    def write(*args, **kwargs):
        receipt = original_write(*args, **kwargs)
        (target / 'two').write_text('concurrent writer')
        return receipt
    monkeypatch.setattr(harness, 'write_bytes', write)
    result = execute(setup, outputs=[{'path': str(target / name), 'name': name, 'expected_sha256': 'missing'}
                                     for name in ('one', 'two')], on_complete=saved.append)
    assert len(result['outputs']) == 1 and result['outputs'] == saved
    assert 'changed' in result['publication_error']
    assert (target / 'two').read_text() == 'concurrent writer'


@pytest.mark.skipif(os.name != 'nt', reason='Windows transfer integration')
def test_cancellation_waits_for_worker_cleanup(setup, monkeypatch):
    entered, cleaned = threading.Event(), threading.Event()
    def child(*a, cancel, **k):
        entered.set()
        assert cancel.wait(5)
        time.sleep(.05)
        cleaned.set()
        return {'exit_code': 1, 'termination_reason': 'cancelled', 'stdout': '', 'stderr': ''}
    monkeypatch.setattr(runner, 'run_isolated', child)
    async def scenario():
        task = asyncio.create_task(runner.execute_python(setup[1], setup[0], 'pass'))
        while not entered.is_set():
            await asyncio.sleep(.01)
        task.cancel()
        await asyncio.sleep(.01)
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert cleaned.is_set()
    asyncio.run(scenario())


@pytest.mark.skipif(os.name != 'nt' or os.environ.get('AWB_TEST_WINDOWS_LPAC') != '1',
                    reason='Set AWB_TEST_WINDOWS_LPAC=1 for real Windows LPAC tests')
def test_real_windows_python_file_network_child_boundary(tmp_path):
    import ctypes
    from ctypes import wintypes as W
    import winreg
    from workbench.windows_sandbox import _Native
    read, write, secret, public = (tmp_path / name for name in ('read', 'write', 'secret', 'lpac-public'))
    for folder in (read, write, secret, public):
        folder.mkdir()
    (read / 'numbers.txt').write_text('9')
    (secret / 'private.txt').write_text('never read')
    (public / 'published.txt').write_text('LPAC public fixture')
    harness = PathHarness([str(read)], [str(write)], [str(secret), str(public)])
    native, public_sid = _Native(), ctypes.c_void_p()
    native.a.ConvertStringSidToSidW.argtypes = [W.LPCWSTR, ctypes.POINTER(ctypes.c_void_p)]
    native.a.ConvertStringSidToSidW.restype = W.BOOL
    assert native.a.ConvertStringSidToSidW('S-1-15-2-2', ctypes.byref(public_sid))
    registry_probe = 'Software\\AgentWorkbenchSandboxProbe.' + uuid.uuid4().hex
    with winreg.CreateKey(winreg.HKEY_CURRENT_USER, registry_probe) as key:
        winreg.SetValueEx(key, 'Secret', 0, winreg.REG_SZ, 'private synthetic value')
    native.acl(public, public_sid)
    config = {'enabled': True, 'executable': str(Path(sys.base_prefix) / 'python.exe'),
              'library_paths': [str(Path(sys.prefix) / 'Lib' / 'site-packages')],
              'timeout_seconds': 30, 'max_memory_mb': 512, 'max_output_bytes': 65536}
    code = f'''
import math,os,socket,subprocess,sys,winreg
from pathlib import Path
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
assert math.sqrt(float(Path('input.txt').read_text()))==3
assert 'AWB_SECRET_TEST' not in os.environ
# Existing Windows LPAC-public resources are an explicit direct-I/O exception,
# even when transfer APIs deny the folder. The harness itself still rejects it.
assert Path({str(public/'published.txt')!r}).read_text()=='LPAC public fixture'
for access in [winreg.KEY_READ,winreg.KEY_SET_VALUE]:
    try: winreg.OpenKey(winreg.HKEY_CURRENT_USER,{registry_probe!r},0,access)
    except OSError: pass
    else: raise AssertionError('Private HKCU registry access escaped')
for path,mode in [({str(secret/'private.txt')!r},'r'),({str(write/'bypass.txt')!r},'w')]:
    try: open(path,mode)
    except PermissionError: pass
    else: raise AssertionError('Host access escaped')
try: subprocess.run([sys.executable,'-c','pass'])
except OSError: pass
else: raise AssertionError('Child creation escaped')
try:
    s=socket.socket();s.settimeout(1);s.connect(('127.0.0.1',8818))
except OSError as error:
    assert error.winerror==10013,repr(error)
else: raise AssertionError('Network escaped')
doc=Document();doc.add_paragraph('LPAC Word');doc.save('sample.docx')
book=Workbook();book.active['A1']=42;book.save('sample.xlsx')
deck=Presentation();deck.slides.add_slide(deck.slide_layouts[6]);deck.save('sample.pptx')
print('隔離OK')
'''
    old = os.environ.get('AWB_SECRET_TEST')
    os.environ['AWB_SECRET_TEST'] = 'never inherit'
    try:
        result = asyncio.run(runner.execute_python(config, harness, code,
                          inputs=[{'path': str(read/'numbers.txt'), 'name': 'input.txt'}],
                          outputs=[{'name': 'sample'+ext, 'path': str(write/('sample'+ext)), 'expected_sha256': 'missing'}
                                   for ext in ('.docx','.xlsx','.pptx')]))
    finally:
        with winreg.OpenKey(winreg.HKEY_CURRENT_USER, registry_probe) as key:
            assert winreg.QueryValueEx(key, 'Secret')[0] == 'private synthetic value'
        winreg.DeleteKey(winreg.HKEY_CURRENT_USER, registry_probe)
        native.acl(public, public_sid, remove=True)
        native.k.LocalFree(public_sid)
        if old is None:
            os.environ.pop('AWB_SECRET_TEST', None)
        else:
            os.environ['AWB_SECRET_TEST'] = old
    assert result['exit_code'] == 0, result
    assert result['termination_reason'] is None, result
    assert '隔離OK' in result['stdout']
    assert 'publication_error' not in result, result.get('publication_error')
    assert len(result['outputs']) == 3, result
    assert not (write/'bypass.txt').exists()


@pytest.mark.skipif(os.name != 'nt' or os.environ.get('AWB_TEST_WINDOWS_LPAC') != '1',
                    reason='Set AWB_TEST_WINDOWS_LPAC=1 for real Windows LPAC tests')
@pytest.mark.parametrize('network_enabled', [False, True])
def test_real_windows_limits_and_job_cancellation(tmp_path, network_enabled):
    from workbench.windows_sandbox import run_isolated
    import ctypes
    root = Path(sys.base_prefix)
    executable = root / 'python.exe'
    access = [(root, False), (root/'Lib', False), (executable, False), (root/'Lib'/'encodings', True)]
    access += [(item, False) for item in root.glob('*.dll')]
    access += [(item, False) for item in (root/'Lib').glob('*.py')]
    access += [(root/'Lib'/name, True) for name in ('collections', 'importlib', 're')]
    def run(code, *, timeout=5, cancel=None, output=1024):
        return run_isolated([str(executable), '-I', '-S', '-B', '-u', '-X', 'utf8', '-c', code],
                            tmp_path, access, timeout=timeout, memory_mb=128, output_limit=output,
                            workspace_limit=1024*1024, cancel=cancel, network_enabled=network_enabled)
    timed = run('while True: pass', timeout=.2)
    assert timed['termination_reason'] == 'timeout', timed
    combined = run("import sys;sys.stdout.write('x'*800);sys.stderr.write('y'*800)")
    assert combined['termination_reason'] == 'output_limit', combined
    assert len(combined['stdout'].encode()) + len(combined['stderr'].encode()) == 1024
    memory = run("try:\n x=bytearray(300*1024*1024)\nexcept MemoryError:\n print('memory denied')")
    assert memory['exit_code'] == 0 and memory['stdout'].strip() == 'memory denied', memory
    stream = run(f"open({str(tmp_path/'stream.txt')!r},'w').write('base');open({str(tmp_path/'stream.txt')+':hidden'!r},'w').write('hidden')")
    assert stream['termination_reason'] == 'workspace_limit', stream
    # Remove the synthetic ADS fixture before the next independent case.
    (tmp_path/'stream.txt').unlink()
    event, started = threading.Event(), tmp_path/'started.txt'
    code = f"import os;open({str(started)!r},'w').write(str(os.getpid()))\nwhile True: pass"
    async def cancellation():
        pending = asyncio.create_task(asyncio.to_thread(run, code, cancel=event))
        deadline = time.monotonic()+20
        while not started.exists() and not pending.done() and time.monotonic()<deadline:
            await asyncio.sleep(.02)
        assert started.exists(), await pending if pending.done() else 'child did not start'
        pid = int(started.read_text())
        event.set()
        result = await pending
        assert result['termination_reason'] == 'cancelled', result
        kernel = ctypes.WinDLL('kernel32', use_last_error=True)
        kernel.OpenProcess.argtypes = [ctypes.c_uint32, ctypes.c_int, ctypes.c_uint32]
        kernel.OpenProcess.restype = ctypes.c_void_p
        handle = kernel.OpenProcess(0x100000, False, pid)
        if handle:
            kernel.WaitForSingleObject.argtypes = [ctypes.c_void_p, ctypes.c_uint32]
            kernel.CloseHandle.argtypes = [ctypes.c_void_p]
            try:
                assert kernel.WaitForSingleObject(handle, 0) == 0
            finally:
                kernel.CloseHandle(handle)
    asyncio.run(cancellation())


@pytest.mark.skipif(os.name != 'nt' or os.environ.get('AWB_TEST_WINDOWS_LPAC') != '1',
                    reason='Set AWB_TEST_WINDOWS_LPAC=1 for real Windows LPAC tests')
def test_real_partial_acl_failure_rolls_back(tmp_path, monkeypatch):
    import ctypes
    from ctypes import wintypes as W
    from workbench.windows_sandbox import _Native, run_isolated
    runtime, workspace = tmp_path/'runtime', tmp_path/'workspace'
    runtime.mkdir(); workspace.mkdir()
    child = runtime/'module.py'
    child.write_text('original')
    native = _Native()
    native.a.ConvertSecurityDescriptorToStringSecurityDescriptorW.argtypes = [ctypes.c_void_p, W.DWORD, W.DWORD,
                                                                             ctypes.POINTER(ctypes.c_void_p), ctypes.POINTER(W.DWORD)]
    native.a.ConvertSecurityDescriptorToStringSecurityDescriptorW.restype = W.BOOL
    def dacl(path):
        acl, descriptor, text, length = ctypes.c_void_p(), ctypes.c_void_p(), ctypes.c_void_p(), W.DWORD()
        assert native.a.GetNamedSecurityInfoW(str(path), 1, 4, None, None, ctypes.byref(acl), None, ctypes.byref(descriptor)) == 0
        try:
            assert native.a.ConvertSecurityDescriptorToStringSecurityDescriptorW(descriptor, 1, 4, ctypes.byref(text), ctypes.byref(length))
            return ctypes.wstring_at(text)
        finally:
            native.k.LocalFree(text);native.k.LocalFree(descriptor)
    before = [dacl(path) for path in (runtime, child)]
    original, calls = _Native.acl, []
    def partial(self, path, sid, **kwargs):
        original(self, path, sid, **kwargs)
        if path == runtime:
            calls.append(kwargs.get('remove', False))
            if not kwargs.get('remove', False):
                raise SandboxError('synthetic partially applied permission failure')
    monkeypatch.setattr(_Native, 'acl', partial)
    with pytest.raises(SandboxError, match='synthetic'):
        run_isolated(['never executed'], workspace, [(runtime, True)], timeout=1, memory_mb=128,
                     output_limit=1024, workspace_limit=1024*1024)
    assert calls == [False, True]
    # SetSecurityInfo may mark an existing DACL auto-inherited (AI); all ACEs
    # and other protection flags must remain identical after removing our SID.
    assert [dacl(path).replace('D:AI', 'D:') for path in (runtime, child)] == [value.replace('D:AI', 'D:') for value in before]


def _minimal_network_runtime():
    root = Path(sys.base_prefix)
    executable = root / 'python.exe'
    access = [(root, False), (root/'Lib', False), (executable, False), (root/'DLLs', True)]
    access += [(item, False) for item in root.glob('*.dll')]
    access += [(item, False) for item in (root/'Lib').glob('*.py')]
    access += [(root/'Lib'/name, True) for name in ('encodings', 'collections', 'importlib', 're', 'ctypes', 'json')]
    return executable, access


_TOKEN_CAPABILITIES_PROBE = r'''
import ctypes,json,socket,subprocess,sys
from ctypes import wintypes as W
kernel=ctypes.WinDLL('kernel32',use_last_error=True)
advapi=ctypes.WinDLL('advapi32',use_last_error=True)
kernel.GetCurrentProcess.restype=W.HANDLE
kernel.CloseHandle.argtypes=[W.HANDLE]
kernel.LocalFree.argtypes=[ctypes.c_void_p]
kernel.LocalFree.restype=ctypes.c_void_p
advapi.OpenProcessToken.argtypes=[W.HANDLE,W.DWORD,ctypes.POINTER(W.HANDLE)]
advapi.GetTokenInformation.argtypes=[W.HANDLE,ctypes.c_int,ctypes.c_void_p,W.DWORD,ctypes.POINTER(W.DWORD)]
advapi.ConvertSidToStringSidW.argtypes=[ctypes.c_void_p,ctypes.POINTER(ctypes.c_void_p)]
class SidAttributes(ctypes.Structure):
    _fields_=[('sid',ctypes.c_void_p),('attributes',W.DWORD)]
class Groups(ctypes.Structure):
    _fields_=[('count',W.DWORD),('groups',SidAttributes*1)]
token=W.HANDLE()
assert advapi.OpenProcessToken(kernel.GetCurrentProcess(),8,ctypes.byref(token)),ctypes.get_last_error()
capabilities=[]
try:
    size=W.DWORD()
    advapi.GetTokenInformation(token,30,None,0,ctypes.byref(size))
    buffer=ctypes.create_string_buffer(size.value)
    assert advapi.GetTokenInformation(token,30,buffer,size,ctypes.byref(size)),ctypes.get_last_error()
    count=W.DWORD.from_buffer(buffer).value
    groups=(SidAttributes*count).from_buffer(buffer,Groups.groups.offset)
    for item in groups:
        text=ctypes.c_void_p()
        assert advapi.ConvertSidToStringSidW(item.sid,ctypes.byref(text)),ctypes.get_last_error()
        try: capabilities.append(ctypes.wstring_at(text))
        finally: kernel.LocalFree(text)
finally: kernel.CloseHandle(token)
try: subprocess.run([sys.executable,'-c','pass'])
except OSError: pass
else: raise AssertionError('Child creation escaped')
print(json.dumps({'capabilities':capabilities}))
'''


@pytest.mark.skipif(os.name != 'nt' or os.environ.get('AWB_TEST_WINDOWS_LPAC') != '1',
                    reason='Set AWB_TEST_WINDOWS_LPAC=1 for real Windows LPAC tests')
@pytest.mark.parametrize('network_enabled', [False, True])
def test_real_windows_network_token_and_loopback_boundary(tmp_path, network_enabled):
    import socket
    from workbench.windows_sandbox import run_isolated
    executable, access = _minimal_network_runtime()
    private = tmp_path/'private.txt'
    private.write_text('synthetic secret')
    workspace = tmp_path/'workspace'
    workspace.mkdir()
    # A listening socket proves denial is policy, not a closed port. No bytes
    # are exchanged. Windows normally blocks same-host AppContainer sockets.
    with socket.socket() as listener:
        listener.bind(('127.0.0.1', 0))
        listener.listen(1)
        port = listener.getsockname()[1]
        code = _TOKEN_CAPABILITIES_PROBE + f'''
for mode in ('r','w'):
    try: open({str(private)!r},mode)
    except PermissionError: pass
    else: raise AssertionError('Private host access escaped')
try:
    with socket.socket() as probe:
        probe.settimeout(2)
        probe.connect(('127.0.0.1',{port}))
except OSError as error:
    # Depending on Windows filtering, loopback can be refused immediately or
    # silently dropped. With networking disabled the socket is denied outright.
    assert getattr(error,'winerror',None)==10013 or ({network_enabled!r} and isinstance(error,TimeoutError)),repr(error)
else: raise AssertionError('Unexpected loopback exemption')
'''
        result = run_isolated([str(executable), '-I', '-S', '-B', '-u', '-X', 'utf8', '-c', code],
                              workspace, access, timeout=10, memory_mb=128, output_limit=4096,
                              workspace_limit=1024*1024, network_enabled=network_enabled)
        listener.setblocking(False)
        with pytest.raises(BlockingIOError):
            listener.accept()
    assert result['exit_code'] == 0 and result['termination_reason'] is None, result
    caps = set(json.loads(result['stdout'])['capabilities'])
    assert ('S-1-15-3-1' in caps) is network_enabled
    assert ('S-1-15-3-3' in caps) is network_enabled
    assert 'S-1-15-3-2' not in caps  # No public-network inbound capability.
    assert result['network'] == ('enabled' if network_enabled else 'disabled')
    assert result['network_capabilities'] == (['internetClient', 'privateNetworkClientServer'] if network_enabled else [])
    assert private.read_text() == 'synthetic secret'


@pytest.mark.skipif(os.name != 'nt' or os.environ.get('AWB_TEST_WINDOWS_NETWORK') != '1',
                    reason='Set AWB_TEST_WINDOWS_NETWORK=1 for opt-in external TCP connectivity check')
@pytest.mark.parametrize('network_enabled', [False, True])
def test_real_windows_external_tcp_without_payload(tmp_path, network_enabled):
    import socket
    from workbench.windows_sandbox import run_isolated
    executable, access = _minimal_network_runtime()
    # This opt-in check contacts only a fixed public test domain. It sends no
    # application payload, credentials, user files or HTTP requests.
    # Resolve in the parent for OFF so a DNS refusal cannot mask socket policy.
    address = 'example.com' if network_enabled else socket.gethostbyname('example.com')
    code = f'''
import socket
try:
    with socket.create_connection(({address!r},443),timeout=8): pass
except OSError as error:
    if {network_enabled!r}: raise
    assert error.winerror==10013,repr(error)
    print('network denied')
else:
    assert {network_enabled!r},'Network escaped'
    print('connected without payload')
'''
    result = run_isolated([str(executable), '-I', '-S', '-B', '-u', '-X', 'utf8', '-c', code],
                          tmp_path, access, timeout=12, memory_mb=128, output_limit=1024,
                          workspace_limit=1024*1024, network_enabled=network_enabled)
    assert result['exit_code'] == 0 and result['termination_reason'] is None, result
    assert result['stdout'].strip() == ('connected without payload' if network_enabled else 'network denied')
