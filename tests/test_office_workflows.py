"""Exercise real Office files through native API tool calls, without Office automation."""
import hashlib
import json
import zipfile

import pytest
from aiohttp import web
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

from workbench.engine import Engine
from workbench.harness import HarnessError, PathHarness, ToolExecutor
from test_engine import settled, start
from test_engine_integration import api_server, configured_settings
from test_tool_result_integrity import WIRE_REPLY


def returned(kind, body, ident):
    """Find a native tool result by ID across multiple assistant turns."""
    if kind == "local":
        value = next(item["content"] for item in body["messages"] if item.get("tool_call_id") == ident)
    elif kind == "openai":
        value = next(item["output"] for item in body["input"]
                     if item.get("type") == "function_call_output" and item["call_id"] == ident)
    else:
        value = next(block["content"] for item in body["messages"] if isinstance(item["content"], list)
                     for block in item["content"]
                     if block.get("type") == "tool_result" and block["tool_use_id"] == ident)
    result = json.loads(value)
    assert result.get("ok") is not False, result
    return result


def initial_content(extension):
    return {
        "docx": {"paragraphs": ["初稿", "保持する本文"], "tables": [[["項目", "初期値"]]]},
        "xlsx": {"cells": [{"sheet": "Sheet", "cell": "A1", "value": "初稿"},
                            {"sheet": "Sheet", "cell": "B2", "value": 42}]},
        "pptx": {"slides": [{"title": "初稿", "body": ["保持する本文"]}]},
    }[extension]


def edit_content(extension):
    return {
        "docx": {"operations": [{"type": "replace_paragraph", "index": 0, "text": "改訂版"},
                                {"type": "replace_cell", "table": 0, "row": 0, "column": 1, "text": "更新値"}]},
        "xlsx": {"cells": [{"sheet": "Sheet", "cell": "A1", "value": "改訂版"},
                            {"sheet": "Sheet", "cell": "C2", "value": "=SUM(B2:B2)"}], "allow_formulas": True},
        "pptx": {"operations": [{"slide": 0, "shape": 0, "paragraph": 0, "text": "改訂版"}]},
    }[extension]


def edit_name(extension):
    return "xlsx_write" if extension == "xlsx" else extension + "_edit"


def assert_contents(values, *, edited=False):
    title = "改訂版" if edited else "初稿"
    assert [p["text"] for p in values["docx"]["paragraphs"]] == [title, "保持する本文"]
    assert values["docx"]["tables"] == [[["項目", "更新値" if edited else "初期値"]]]
    assert values["xlsx"]["cells"][:2] == [{"cell": "A1", "value": title, "formula": False},
                                         {"cell": "B2", "value": 42, "formula": False}]
    if edited:
        assert values["xlsx"]["cells"][2] == {"cell": "C2", "value": "=SUM(B2:B2)", "formula": True}
    assert values["pptx"]["slides"][0]["shapes"][0]["paragraphs"] == [title]
    assert values["pptx"]["slides"][0]["shapes"][1]["paragraphs"] == ["保持する本文"]


@pytest.mark.parametrize("kind", WIRE_REPLY)
async def test_native_office_create_read_edit_read_and_local_search(tmp_path, kind):
    requests, errors, original_hashes = [], [], {}
    paths = {ext: "office/reports/報告書." + ext for ext in ("docx", "xlsx", "pptx")}

    async def handler(request):
        try:
            body = await request.json()
            requests.append(body)
            step = len(requests)
            if step == 1:
                assert "search_paths" in json.dumps(body["tools"])
                calls = [("reserve", "reserve_paths", {"paths": ["office"]}),
                         ("mkdir", "create_directory", {"path": "office/reports"})]
                calls.extend(("write-" + ext, ext + "_write", {"path": path, "expected_sha256": "missing",
                              **initial_content(ext)}) for ext, path in paths.items())
                calls.extend(("read-" + ext, ext + "_read", {"path": path}) for ext, path in paths.items())
                return WIRE_REPLY[kind](calls)
            if step == 2:
                assert returned(kind, body, "mkdir")["operation"] == "created"
                values = {ext: returned(kind, body, "read-" + ext) for ext in paths}
                assert_contents(values)
                for ext in paths:
                    original_hashes[ext] = values[ext]["sha256"]
                    assert returned(kind, body, "write-" + ext)["sha256"] == original_hashes[ext]
                calls = [("edit-" + ext, edit_name(ext), {"path": path,
                          "expected_sha256": original_hashes[ext], **edit_content(ext)}) for ext, path in paths.items()]
                calls.extend(("verify-" + ext, ext + "_read", {"path": path}) for ext, path in paths.items())
                calls.extend([("find-paths", "search_paths", {"path": "office", "query": "報告書", "kind": "file"}),
                              ("find-text", "search_text", {"path": "notes", "query": "検索対象"}),
                              ("release", "release_paths", {"paths": ["office"]})])
                return WIRE_REPLY[kind](calls)
            assert step == 3
            values = {ext: returned(kind, body, "verify-" + ext) for ext in paths}
            assert_contents(values, edited=True)
            for ext in paths:
                assert values[ext]["sha256"] != original_hashes[ext]
                assert returned(kind, body, "edit-" + ext)["sha256"] == values[ext]["sha256"]
            found = returned(kind, body, "find-paths")
            assert {item["name"] for item in found["matches"]} == {"報告書." + ext for ext in paths}
            assert not found["truncated"]
            found_text = returned(kind, body, "find-text")
            assert len(found_text["matches"]) == 1 and found_text["matches"][0]["line"] == 2
            returned(kind, body, "release")
            return WIRE_REPLY[kind]([("finish", "finish_work", {"summary": "Created, edited and verified all Office files."})])
        except Exception as exc:
            errors.append(repr(exc))
            return web.json_response({"error": "Office workflow fixture failed"}, status=500)

    async with api_server(handler) as endpoint:
        settings, work = configured_settings(tmp_path, endpoint)
        (work / "notes").mkdir()
        (work / "notes/index.txt").write_text("索引\n検索対象", encoding="utf-8")
        engine = Engine(settings)
        try:
            run, agent = await start(engine, pm_profile=kind, max_workers=0)
            await settled(engine)
            assert not errors, errors
            assert len(requests) == 3 and run["status"] == "done" and agent.status == "done"
            assert not agent.last_error and not engine.reservations
            assert len(agent.output_receipts) == 6
            assert {item["tool"] for item in agent.output_receipts} == {
                "docx_write", "docx_edit", "xlsx_write", "pptx_write", "pptx_edit"}
            document = Document(work / paths["docx"])
            assert [p.text for p in document.paragraphs] == ["改訂版", "保持する本文"]
            assert document.tables[0].cell(0, 1).text == "更新値"
            workbook = load_workbook(work / paths["xlsx"], data_only=False)
            try:
                assert workbook.active["A1"].value == "改訂版"
                assert workbook.active["B2"].value == 42
                assert workbook.active["C2"].value == "=SUM(B2:B2)"
            finally:
                workbook.close()
            slide = Presentation(work / paths["pptx"]).slides[0]
            assert slide.shapes.title.text == "改訂版" and slide.placeholders[1].text == "保持する本文"
        finally:
            await engine.close()


@pytest.mark.parametrize("extension", ("docx", "xlsx", "pptx"))
async def test_office_edits_require_read_write_scope_and_current_hash(tmp_path, extension):
    allowed, outside = tmp_path / "allowed", tmp_path / "outside"
    allowed.mkdir()
    outside.mkdir()
    path = allowed / ("report." + extension)
    tools = ToolExecutor(PathHarness([str(allowed)], [str(allowed)]))
    saved = await tools.execute(extension + "_write", {
        "path": str(path), "expected_sha256": "missing", **initial_content(extension)})
    original = path.read_bytes()
    args = {"path": str(path), "expected_sha256": saved["sha256"], **edit_content(extension)}
    for harness in (PathHarness([], [str(allowed)]), PathHarness([str(allowed)], []),
                    PathHarness([str(allowed)], [str(allowed)], [str(path)])):
        with pytest.raises(HarnessError):
            await ToolExecutor(harness).execute(edit_name(extension), args)
        assert path.read_bytes() == original
    with pytest.raises(HarnessError):
        await tools.execute(edit_name(extension), {**args, "expected_sha256": "0" * 64})
    with pytest.raises(HarnessError):
        await tools.execute(extension + "_write", {"path": str(outside / path.name),
                            "expected_sha256": "missing", **initial_content(extension)})
    assert path.read_bytes() == original and not list(outside.iterdir())


@pytest.mark.parametrize("extension,part", (("docx", "word"), ("xlsx", "xl"), ("pptx", "ppt")))
async def test_renamed_macro_package_is_rejected_before_library_parsing(tmp_path, extension, part, monkeypatch):
    path = tmp_path / ("macro-disguised." + extension)
    tools = ToolExecutor(PathHarness([str(tmp_path)], [str(tmp_path)]))
    await tools.execute(extension + "_write", {"path": str(path), "expected_sha256": "missing",
                                               **initial_content(extension)})
    # An inert marker tests package detection without embedding executable VBA.
    with zipfile.ZipFile(path, "a") as package:
        package.writestr(part + "/vbaProject.bin", b"inert test marker")
    original = path.read_bytes()
    module, parser = {"docx": ("docx", "Document"), "xlsx": ("openpyxl", "load_workbook"),
                      "pptx": ("pptx", "Presentation")}[extension]
    monkeypatch.setattr(module + "." + parser, lambda *_a, **_k: pytest.fail("Unsafe package reached Office parser"))
    with pytest.raises(HarnessError, match="Macros"):
        await tools.execute(extension + "_read", {"path": str(path)})
    with pytest.raises(HarnessError, match="Macros"):
        await tools.execute(edit_name(extension), {"path": str(path),
                            "expected_sha256": hashlib.sha256(original).hexdigest(), **edit_content(extension)})
    assert path.read_bytes() == original
