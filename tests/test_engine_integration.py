"""Real Engine + HTTP provider adapters + Office tools; all APIs are local fixtures.

No UI mock, real model, external service, CLI, or paid inference is used.
"""
import asyncio
import json
import sys
import threading
import time
from contextlib import asynccontextmanager
from types import SimpleNamespace

import pytest
from aiohttp import web

from workbench.config import Settings, validate_settings
from workbench.engine import Engine
from workbench.harness import ToolExecutor


@asynccontextmanager
async def api_server(handler):
    app = web.Application()
    app.router.add_post("/{path:.*}", handler)
    runner = web.AppRunner(app)
    await runner.setup()
    site = web.TCPSite(runner, "127.0.0.1", 0)
    await site.start()
    try:
        yield f"http://127.0.0.1:{site._server.sockets[0].getsockname()[1]}/v1"
    finally:
        await runner.cleanup()


COUNTERS = ('model_calls', 'tool_calls', 'auto_collaborations')


class RunProgressGuard:
    """Observed progress watchdog, bounded by the run's existing remaining budget.

    An observed counter increase has no precise occurrence timestamp. This is
    not a 15-second performance SLO; the fixed monotonic ceiling never rearms.
    """
    def __init__(self, runs, run, monotonic, wall, timeout=15):
        self.runs, self.run, self.clock = runs, run, monotonic
        self.run_id, self.config = run['id'], run['_config']
        self.created_at = run['created_at']
        limits = self.config['limits']
        self.caps = tuple(limits['max_' + name] for name in COUNTERS)
        self.run_seconds = limits['max_run_seconds']
        if any(type(cap) is not int or cap < 0 for cap in self.caps):
            raise AssertionError('invalid_limits')
        if type(self.run_seconds) is not int or self.run_seconds <= 0:
            raise AssertionError('invalid_run_budget')
        start = monotonic()
        # Same wall-clock origin as Engine._run_deadline; a single translation.
        remaining = self.run_seconds - max(0, wall() - self.created_at)
        self.hard_deadline = start + max(0, remaining)
        self.timeout, self.last_at = timeout, start
        self.previous = self.values()

    def values(self):
        run = self.run
        if self.runs.get(self.run_id) is not run or run['id'] != self.run_id:
            raise AssertionError('run_identity_changed')
        if (run['_config'] is not self.config or run['created_at'] != self.created_at
                or tuple(self.config['limits']['max_' + name] for name in COUNTERS) != self.caps
                or self.config['limits']['max_run_seconds'] != self.run_seconds
                or run['max_auto_collaborations'] != self.caps[2]):
            raise AssertionError('run_limits_changed')
        values = tuple(run[name] for name in COUNTERS)
        if any(type(value) is not int or value < 0 for value in values):
            raise AssertionError('invalid_counter')
        if any(value > cap for value, cap in zip(values, self.caps)):
            raise AssertionError('counter_over_budget')
        return values

    def check(self):
        values, now = self.values(), self.clock()
        if any(value < previous for value, previous in zip(values, self.previous)):
            raise AssertionError('counter_regressed')
        # Validate even at completion; a new counter cannot rescue this deadline.
        if now >= self.hard_deadline:
            raise AssertionError('run_deadline')
        if values != self.previous:
            self.last_at, self.previous = now, values
        if now - self.last_at > self.timeout:
            raise AssertionError('no_progress')


async def eventually(predicate, engine=None, timeout=15, *, progress_run=None):
    if progress_run is not None:
        guard = RunProgressGuard(engine.runs, progress_run,
                                 asyncio.get_running_loop().time, time.time, timeout)
        while True:
            try:
                guard.check()
            except AssertionError as error:
                pytest.fail("Integration did not settle (" + str(error) + "): "
                            + json.dumps(engine.snapshot(), ensure_ascii=False))
            if predicate():
                return
            await asyncio.sleep(.01)
    deadline = asyncio.get_running_loop().time() + timeout
    while not predicate():
        if asyncio.get_running_loop().time() > deadline:
            pytest.fail("Integration did not settle: " + json.dumps(engine.snapshot() if engine else {}, ensure_ascii=False))
        await asyncio.sleep(.01)


def configured_settings(tmp_path, endpoint):
    workspace = tmp_path / "documents"
    workspace.mkdir()
    settings = Settings(tmp_path / "state")
    config = settings.value
    for profile in config["providers"]:
        profile["base_url"] = endpoint
        profile["model"] = {"openai": "responses-fixture", "local": "qwen-fixture", "anthropic": "messages-fixture"}[profile["id"]]
    config["paths"] = {"read_roots": [str(workspace)], "write_roots": [str(workspace)], "deny_roots": []}
    config["local"].update(max_concurrent_requests=1, max_retries=0, gpu_guard_enabled=False)
    settings.value = validate_settings(config)
    settings.secrets.update(openai="test-openai-secret", anthropic="test-anthropic-secret")
    return settings, workspace


def openai_reply(calls=(), text=""):
    output = [{"type": "function_call", "call_id": ident, "id": "fc-" + ident, "name": name,
               "arguments": json.dumps(arguments, ensure_ascii=False)} for ident, name, arguments in calls]
    if text:
        output.append({"type": "message", "role": "assistant", "content": [{"type": "output_text", "text": text}]})
    return web.json_response({"status": "completed", "output": output})


def local_reply(calls=(), text=""):
    message = {"role": "assistant", "content": text,
               "tool_calls": [{"id": ident, "type": "function", "function": {"name": name, "arguments": json.dumps(arguments, ensure_ascii=False)}} for ident, name, arguments in calls]}
    return web.json_response({"choices": [{"message": message, "finish_reason": "tool_calls" if calls else "stop"}]})


def anthropic_reply(calls=(), text=""):
    blocks = [{"type": "tool_use", "id": ident, "name": name, "input": arguments} for ident, name, arguments in calls]
    if text:
        blocks.insert(0, {"type": "text", "text": text})
    return web.json_response({"content": blocks, "stop_reason": "tool_use" if calls else "end_turn"})


@pytest.mark.asyncio
async def test_real_engine_mixed_api_team_creates_verifies_office_and_resumes_pm(tmp_path):
    engine = None
    requests = {"pm": [], "docx": [], "pptx": [], "anthropic": []}
    errors = []
    local_started = asyncio.Event()
    local_active = 0
    local_peak = 0
    cloud_overlapped_local = False
    documents = tmp_path / "documents"
    paths = {kind: str(documents / ("result." + kind)) for kind in ("docx", "xlsx", "pptx")}

    async def handler(request):
        nonlocal local_active, local_peak, cloud_overlapped_local
        try:
            body = await request.json()
            if request.path == "/v1/responses":
                assert request.headers["Authorization"] == "Bearer test-openai-secret"
                requests["pm"].append(body)
                if len(requests["pm"]) == 1:
                    return openai_reply([
                        ("spawn-docx", "spawn_worker", {"profile_id": "local", "role": "Document writer", "task": "DOCX worker: create and verify " + paths["docx"]}),
                        ("spawn-pptx", "spawn_worker", {"profile_id": "local", "role": "Slides writer", "task": "PPTX worker: create and verify " + paths["pptx"]}),
                        ("spawn-xlsx", "spawn_worker", {"profile_id": "anthropic", "role": "Spreadsheet writer", "task": "Create and verify " + paths["xlsx"]}),
                    ])
                notifications = []
                for item in body["input"]:
                    if item.get("role") == "user" and isinstance(item.get("content"), str) and item["content"].startswith("Peer data;"):
                        envelope = json.loads(item["content"].split("\n", 1)[1])
                        if envelope["kind"] == "worker_completed":
                            notifications.append(envelope)
                if len({item["sender"] for item in notifications}) == 3:
                    assert all("Verified" in item["body"] and item["untrusted"] for item in notifications)
                    return openai_reply([("pm-complete", "finish_work", {"summary": "All three Office files were created and verified."})])
                return openai_reply(text="Waiting for the independent workers' completion messages.")
            if request.path == "/v1/chat/completions":
                marker = next(item["content"] for item in body["messages"] if item["role"] == "user")
                kind = "docx" if "DOCX worker" in marker else "pptx"
                requests[kind].append(body)
                step = len(requests[kind])
                local_active += 1
                local_peak = max(local_peak, local_active)
                local_started.set()
                try:
                    await asyncio.sleep(.04)
                    assert all(tool["function"]["name"] != "spawn_worker" for tool in body["tools"])
                    if step == 1:
                        arguments = {"path": paths[kind], "expected_sha256": "missing"}
                        if kind == "docx":
                            arguments.update(paragraphs=["Actual engine integration", "日本語の本文"], tables=[[["Checked", "Yes"]]])
                        else:
                            arguments.update(slides=[{"title": "Actual integration", "body": ["Verified through file tools"]}])
                        return local_reply([(kind + "-reserve", "reserve_paths", {"paths": [paths[kind]]}),
                                            (kind + "-write", kind + "_write", arguments)],
                                           "<think>Select the permitted Office tool.</think>\nCreating the assigned file.")
                    if step == 2:
                        write_result = json.loads(next(item["content"] for item in body["messages"] if item.get("tool_call_id") == kind + "-write"))
                        assert write_result.get("ok") is not False and len(write_result["sha256"]) == 64
                        return local_reply([(kind + "-read", kind + "_read", {"path": paths[kind]})])
                    read_result = json.loads(next(item["content"] for item in body["messages"] if item.get("tool_call_id") == kind + "-read"))
                    assert read_result.get("ok") is not False
                    assert (read_result["paragraphs"][1]["text"] == "日本語の本文") if kind == "docx" else (read_result["slides"][0]["shapes"][0]["paragraphs"] == ["Actual integration"])
                    return local_reply([(kind + "-finish", "finish_work", {"summary": "Verified " + kind.upper() + " content via read tool."})])
                finally:
                    local_active -= 1
            assert request.path == "/v1/messages"
            assert request.headers["x-api-key"] == "test-anthropic-secret"
            requests["anthropic"].append(body)
            step = len(requests["anthropic"])
            if step == 1:
                await local_started.wait()
                cloud_overlapped_local = local_active > 0
                return anthropic_reply([
                    ("xlsx-reserve", "reserve_paths", {"paths": [paths["xlsx"]]}),
                    ("xlsx-write", "xlsx_write", {"path": paths["xlsx"], "expected_sha256": "missing", "cells": [
                        {"sheet": "Sheet", "cell": "A1", "value": "Verified"}, {"sheet": "Sheet", "cell": "B2", "value": 42}]})])
            results = [block for message in body["messages"] for block in message.get("content", []) if isinstance(block, dict) and block.get("type") == "tool_result"]
            if step == 2:
                saved = json.loads(next(item["content"] for item in results if item["tool_use_id"] == "xlsx-write"))
                assert saved.get("ok") is not False and len(saved["sha256"]) == 64
                return anthropic_reply([("xlsx-read", "xlsx_read", {"path": paths["xlsx"], "range": "A1:B2"})])
            read = json.loads(next(item["content"] for item in results if item["tool_use_id"] == "xlsx-read"))
            assert read["cells"][-1]["value"] == 42
            return anthropic_reply([("xlsx-finish", "finish_work", {"summary": "Verified XLSX content via read tool."})])
        except Exception as exc:
            errors.append(repr(exc))
            return web.json_response({"error": "fixture failed"}, status=500)

    async with api_server(handler) as endpoint:
        settings, workspace = configured_settings(tmp_path, endpoint)
        engine = Engine(settings)
        try:
            started = await engine.start_run({"task": "Create verified Office files with independent workers.", "pm_profile": "openai", "worker_profiles": ["local", "anthropic"], "max_workers": 3})
            run = engine.runs[started["run"]["id"]]
            await eventually(lambda: run["status"] in {"done", "waiting"}, engine, progress_run=run)
            assert not errors
            assert run["status"] == "done", engine.snapshot()
            assert len(run["agent_ids"]) == 4
            assert [agent.profile_id for agent in engine.agents.values()].count("local") == 2
            assert local_peak == 1 and cloud_overlapped_local
            assert run["max_workers"] == 3 and settings.value["local"]["max_concurrent_requests"] == 1
            assert len(requests["docx"]) == len(requests["pptx"]) == len(requests["anthropic"]) == 3
            lead = engine.agents[run["agent_ids"][0]]
            assert lead.turns >= 2 and lead.status == "done"
            assert any(item["kind"] == "result" and "All three" in item["text"] for item in lead.logs)
            assert sum(item["kind"] == "mail" for item in engine.events) == 3
            assert any(log["thinking"] for agent in engine.agents.values() for log in agent.logs)
            assert engine.reservations == {}
            assert "test-openai-secret" not in json.dumps(engine.snapshot()) and "test-anthropic-secret" not in json.dumps(engine.snapshot())
            from docx import Document
            from openpyxl import load_workbook
            from pptx import Presentation
            assert Document(paths["docx"]).paragraphs[1].text == "日本語の本文"
            workbook = load_workbook(paths["xlsx"])
            assert workbook.active["B2"].value == 42
            workbook.close()
            assert Presentation(paths["pptx"]).slides[0].shapes.title.text == "Actual integration"
            assert set(item.suffix for item in workspace.iterdir()) == {".docx", ".xlsx", ".pptx"}
        finally:
            await engine.close()


@pytest.mark.asyncio
async def test_real_engine_stop_drains_file_write_before_releasing_reservation(tmp_path, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    path = tmp_path / "documents" / "settled.txt"
    requests = []
    original = ToolExecutor._execute
    def delayed(self, name, args):
        if name == "write_text":
            entered.set()
            assert release.wait(5)
        return original(self, name, args)
    monkeypatch.setattr(ToolExecutor, "_execute", delayed)
    async def handler(request):
        requests.append(await request.json())
        return local_reply([("reserve", "reserve_paths", {"paths": [str(path)]}),
                            ("write", "write_text", {"path": str(path), "text": "Settled before stopped", "expected_sha256": "missing"})])
    async with api_server(handler) as endpoint:
        settings, _ = configured_settings(tmp_path, endpoint)
        engine = Engine(settings)
        stopping = None
        try:
            started = await engine.start_run({"task": "Write one file", "pm_profile": "local", "worker_profiles": ["local"], "max_workers": 0})
            run = engine.runs[started["run"]["id"]]
            assert await asyncio.to_thread(entered.wait, 3)
            stopping = asyncio.create_task(engine.stop_run(run["id"]))
            await eventually(lambda: run["status"] == "stopping", engine)
            await asyncio.sleep(.02)
            assert not stopping.done() and engine.lock.locked() and engine.reservations
            assert not path.exists()
            release.set()
            await stopping
            assert run["status"] == "stopped" and engine.reservations == {} and not engine.lock.locked()
            assert path.read_text(encoding="utf-8") == "Settled before stopped"
            assert len(requests) == 1
            pm = engine.agents[run['agent_ids'][0]]
            assert len(pm.output_receipts) == 1 and pm.output_receipts[0]['path'] == str(path)
            assert not pm.results
        finally:
            release.set()
            if stopping:
                await asyncio.gather(stopping, return_exceptions=True)
            await engine.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('case,opt_in,frames,used,expected', [
    ('progress', True, [(8, {'model_calls': 1}, False), (16, {'tool_calls': 1}, False),
                       (24, {'auto_collaborations': 1}, False), (25, {}, True)], 0, None),
    ('default_unchanged', False, [(8, {'model_calls': 1}, False), (16, {'tool_calls': 1}, False)], 0, 'Integration did not settle:'),
    ('stall', True, [(16, {}, False)], 0, 'no_progress'),
    ('time_only', True, [(8, {'observed_at': 8}, False), (16, {'observed_at': 16}, False)], 0, 'no_progress'),
    ('other_run', True, [(8, {'_other': 1}, False), (16, {'_other': 2}, False)], 0, 'no_progress'),
    ('regressed', True, [(8, {'model_calls': 1}, False), (9, {'model_calls': 0}, True)], 0, 'counter_regressed'),
    ('over_budget', True, [(8, {'model_calls': 25}, True)], 0, 'counter_over_budget'),
    ('bool', True, [(8, {'tool_calls': True}, True)], 0, 'invalid_counter'),
    ('limit_changed', True, [(8, {'_cap': 25}, True)], 0, 'run_limits_changed'),
    ('identity_changed', True, [(8, {'_replace': True}, True)], 0, 'run_identity_changed'),
    ('ceiling', True, [(8, {'model_calls': 1}, False), (16, {'tool_calls': 1}, False),
                      (24, {'auto_collaborations': 1}, False), (30, {'model_calls': 2}, True)], 0, 'run_deadline'),
    ('already_used', True, [(8, {'model_calls': 1}, False), (16, {'tool_calls': 1}, False),
                           (20, {'model_calls': 2}, True)], 11, 'run_deadline'),
    ('late_progress_is_observed_not_timed', True, [(16, {'model_calls': 1}, True)], 0, None),
], ids=lambda value: value if isinstance(value, str) else None)
async def test_whole_run_progress_guard_opt_in_preserves_bounds(monkeypatch, case, opt_in, frames, used, expected):
    """Drive the actual helper with finite clock steps, without Engine or sleep."""
    clock, complete, pending = {'now': 0.0}, {'value': False}, list(frames)
    limits = dict(max_model_calls=24, max_tool_calls=60,
                  max_auto_collaborations=24, max_run_seconds=30)
    run = dict(id='target', created_at=1000.0 - used, _config={'limits': limits},
               max_auto_collaborations=24, model_calls=0, tool_calls=0, auto_collaborations=0)
    runs = {'target': run, 'other': {'model_calls': 0}}
    engine = SimpleNamespace(snapshot=lambda: {'target': run['id'], 'at': clock['now']})
    if opt_in:
        engine.runs = runs

    async def advance(delay):
        assert delay == .01 and pending, 'Guard exhausted the finite observation schedule'
        at, changes, done = pending.pop(0)
        clock['now'], complete['value'] = at, done
        changes = dict(changes)
        if '_other' in changes:
            runs['other']['model_calls'] = changes.pop('_other')
        if '_cap' in changes:
            limits['max_model_calls'] = changes.pop('_cap')
        if '_replace' in changes:
            changes.pop('_replace')
            runs['target'] = dict(run)
        run.update(changes)

    # Replace only this module's clock facade, never pytest's event-loop clock.
    monkeypatch.setattr(sys.modules[__name__], 'asyncio', SimpleNamespace(
        get_running_loop=lambda: SimpleNamespace(time=lambda: clock['now']), sleep=advance))
    monkeypatch.setattr(sys.modules[__name__], 'time', SimpleNamespace(time=lambda: 1000.0))
    kwargs = {'progress_run': run} if opt_in else {}
    if expected:
        with pytest.raises(pytest.fail.Exception, match=expected):
            await eventually(lambda: complete['value'], engine, **kwargs)
    else:
        await eventually(lambda: complete['value'], engine, **kwargs)
        assert complete['value']
    assert not pending
